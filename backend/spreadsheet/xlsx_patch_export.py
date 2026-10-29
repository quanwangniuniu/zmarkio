"""Patch a copy of an imported .xlsx from a Phase 3 diff.

This module does not rebuild a workbook and does not catch its own failures.
The export view decides whether to fall back to build_sheet_workbook.
"""

import io
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from .models import Cell, SheetColumn, SheetRow
from .sparkline import is_sparkline
from .xlsx_export import (
    _EXCEL_FONT_NAMES,
    _apply_fill,
    _cell_export_value,
    _excel_number_format,
    _excel_rgb,
)

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange

# _excel_font omits name, size, and color when the web value is null.
# A brand-new openpyxl cell then uses these library defaults. This repo does
# not define its own Excel font name, point size, or cleared text color.
# The grid draws a null font size at 12px; the Excel exporter does not write 12.
_OPENPYXL_DEFAULT_FONT_NAME = 'Calibri'
_OPENPYXL_DEFAULT_FONT_SIZE = 11
_OPENPYXL_DEFAULT_FONT_COLOR = Color(theme=1)

_FONT_FIELDS = (
    'bold',
    'italic',
    'strikethrough',
    'text_color',
    'font_family',
    'font_size',
)


class XlsxPatchExportError(Exception):
    """The original workbook cannot be patched safely."""


def build_patched_workbook(source, diff, *, workbook_bytes=None, worksheet_name=None) -> bytes:
    """Return xlsx bytes for an original workbook with diff applied to one worksheet."""
    if workbook_bytes is None:
        if not getattr(source, 'file', None):
            raise XlsxPatchExportError('The original Excel file is missing.')
        try:
            with source.file.open('rb') as handle:
                workbook = load_workbook(handle)
        except XlsxPatchExportError:
            raise
        except Exception as exc:
            raise XlsxPatchExportError('The original Excel file could not be opened.') from exc
    else:
        if not isinstance(workbook_bytes, (bytes, bytearray)) or not workbook_bytes:
            raise XlsxPatchExportError('The original Excel file is missing.')
        try:
            workbook = load_workbook(io.BytesIO(workbook_bytes))
        except Exception as exc:
            raise XlsxPatchExportError('The original Excel file could not be opened.') from exc

    selected_name = worksheet_name if isinstance(worksheet_name, str) and worksheet_name.strip() else getattr(source, 'worksheet_name', '')
    sheet = source.sheet
    try:
        worksheet = _worksheet(workbook, selected_name)
        _replay_structure(worksheet, diff.get('structural_operations') or [])
        _reject_sparkline_changes(sheet, diff.get('cells') or {})
        for key, change in (diff.get('cells') or {}).items():
            _patch_value(worksheet, sheet, key, change)
        for key, change in (diff.get('formats') or {}).items():
            _patch_format(worksheet, key, change or {})
        highlights = diff.get('highlights') or {}
        for key, color in (highlights.get('cells') or {}).items():
            _patch_highlight(worksheet, key, color)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()

def _worksheet(workbook, worksheet_name):
    name = (worksheet_name or '').strip()
    if not name or name not in workbook.sheetnames:
        raise XlsxPatchExportError(
            f'Worksheet {worksheet_name!r} is not in the original workbook.'
        )
    return workbook[name]

def _replay_structure(worksheet, operations):
    """Apply validated insert and delete operations before cell patches."""
    if not operations:
        return
    _reject_unshifted_formulas(worksheet)
    for operation in operations:
        if not isinstance(operation, dict):
            raise XlsxPatchExportError('Bad structural operation.')
        op_type = operation.get('op_type')
        anchor = operation.get('anchor_position')
        count = operation.get('count')
        if isinstance(anchor, bool) or not isinstance(anchor, int) or anchor < 0:
            raise XlsxPatchExportError('Bad structural anchor.')
        if isinstance(count, bool) or not isinstance(count, int) or count < 1:
            raise XlsxPatchExportError('Bad structural count.')
        excel_index = anchor + 1
        if op_type in ('ROW_INSERT', 'ROW_DELETE'):
            axis = 'row'
        elif op_type in ('COL_INSERT', 'COL_DELETE'):
            axis = 'column'
        else:
            raise XlsxPatchExportError(f'Unsupported structural operation {op_type!r}.')
        
        deleting = op_type.endswith('_DELETE')
        _reject_unsupported_structure(worksheet, axis, excel_index, count, deleting)
        if deleting:
            _drop_deleted_metadata(worksheet, axis, excel_index, count)
        moved_merges = _take_merges_after_edit(worksheet, axis, excel_index, count, deleting)
        _shift_dimensions_after_edit(worksheet, axis, excel_index, count, deleting)
        if op_type == 'ROW_INSERT':
            worksheet.insert_rows(excel_index, count)
        elif op_type == 'ROW_DELETE':
            worksheet.delete_rows(excel_index, count)
        elif op_type == 'COL_INSERT':
            worksheet.insert_cols(excel_index, count)
        else:
            worksheet.delete_cols(excel_index, count)
        for ref in moved_merges:
            worksheet.merge_cells(ref)

def _reject_unshifted_formulas(worksheet):
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            text = value if isinstance(value, str) else getattr(value, 'text', None)
            if isinstance(text, str) and text.startswith('=') and ('$' in text or '!' in text):
                raise XlsxPatchExportError('This formula reference cannot be shifted.')

def _reject_unsupported_structure(worksheet, axis, excel_index, count, deleting):
    """Refuse features an insert or delete would leave on the old index."""
    if worksheet.conditional_formatting._cf_rules:
        raise XlsxPatchExportError('Conditional formatting cannot be shifted.')
    if worksheet.data_validations.dataValidation:
        raise XlsxPatchExportError('Data validation cannot be shifted.')
    if worksheet.tables:
        raise XlsxPatchExportError('Tables cannot be shifted.')
    if worksheet._charts or worksheet._images:
        raise XlsxPatchExportError('Charts or images cannot be shifted.')
    if worksheet.freeze_panes:
        raise XlsxPatchExportError('Frozen panes cannot be shifted.')
    if worksheet.auto_filter.ref:
        raise XlsxPatchExportError('An auto filter cannot be shifted.')
    if _defined_name_uses_sheet(worksheet):
        raise XlsxPatchExportError('Named ranges cannot be shifted.')
    span_end = excel_index + count
    for merged in worksheet.merged_cells.ranges:
        if axis == 'row' and _feature_blocks_structure(
            merged.min_row, merged.max_row, excel_index, span_end, deleting,
        ):
            raise XlsxPatchExportError('A merged range crosses this row change.')
        if axis == 'column' and _feature_blocks_structure(
            merged.min_col, merged.max_col, excel_index, span_end, deleting,
        ):
            raise XlsxPatchExportError('A merged range crosses this column change.')
    if axis == 'row':
        for index, dimension in worksheet.row_dimensions.items():
            if isinstance(index, int) and (dimension.height is not None or dimension.hidden):
                if _feature_blocks_structure(index, index, excel_index, span_end, deleting):
                    raise XlsxPatchExportError('Row heights cannot be shifted.')
        return
    for letter, dimension in worksheet.column_dimensions.items():
        if dimension.width is None and not dimension.hidden:
            continue
        index = column_index_from_string(letter)
        if _feature_blocks_structure(index, index, excel_index, span_end, deleting):
            raise XlsxPatchExportError('Column widths cannot be shifted.')

def _defined_name_uses_sheet(worksheet):
    defined_names = getattr(worksheet.parent, 'defined_names', None)
    if not defined_names:
        return False
    title = worksheet.title
    quoted = "'" + title.replace("'", "''") + "'"
    for defined in defined_names.values():
        text = defined.attr_text or ''
        if quoted in text or f'{title}!' in text:
            return True
    return False

def _feature_blocks_structure(start, end, cut, span_end, deleting):
    """True when this feature crosses the edit and cannot be moved as a whole."""
    if end < cut:
        return False
    if deleting and start >= cut and end < span_end:
        return False
    if start >= (span_end if deleting else cut):
        return False
    return True


def _drop_deleted_metadata(worksheet, axis, excel_index, count):
    """Remove dimensions and merges that occupy only the rows or columns being deleted."""
    span_end = excel_index + count
    if axis == 'row':
        for index in list(worksheet.row_dimensions):
            if isinstance(index, int) and excel_index <= index < span_end:
                del worksheet.row_dimensions[index]
        for merged in list(worksheet.merged_cells.ranges):
            if merged.min_row >= excel_index and merged.max_row < span_end:
                worksheet.unmerge_cells(str(merged))
        return
    for letter in list(worksheet.column_dimensions):
        index = column_index_from_string(letter)
        if excel_index <= index < span_end:
            del worksheet.column_dimensions[letter]
    for merged in list(worksheet.merged_cells.ranges):
        if merged.min_col >= excel_index and merged.max_col < span_end:
            worksheet.unmerge_cells(str(merged))

def _take_merges_after_edit(worksheet, axis, cut, count, deleting):
    span_end = cut + count
    move_from = span_end if deleting else cut
    moved = []
    for merged in list(worksheet.merged_cells.ranges):
        start = merged.min_row if axis == 'row' else merged.min_col
        if start < move_from:
            continue
        worksheet.unmerge_cells(range_string=str(merged))
        shifted = CellRange(str(merged))
        delta = -count if deleting else count
        if axis == 'row':
            shifted.shift(row_shift=delta)
        else:
            shifted.shift(col_shift=delta)
        moved.append(str(shifted))
    return moved


def _shift_dimensions_after_edit(worksheet, axis, cut, count, deleting):
    move_from = cut + count if deleting else cut
    delta = -count if deleting else count
    if axis == 'row':
        pending = []
        for index in list(worksheet.row_dimensions):
            if not isinstance(index, int) or index < move_from:
                continue
            dimension = worksheet.row_dimensions[index]
            if dimension.height is None and not dimension.hidden:
                continue
            pending.append((index, dimension.height, dimension.hidden))
            del worksheet.row_dimensions[index]
        for index, height, hidden in pending:
            target = worksheet.row_dimensions[index + delta]
            target.height = height
            target.hidden = hidden
        return
    pending = []
    for letter in list(worksheet.column_dimensions):
        index = column_index_from_string(letter)
        if index < move_from:
            continue
        dimension = worksheet.column_dimensions[letter]
        if dimension.width is None and not dimension.hidden:
            continue
        pending.append((index, dimension.width, dimension.hidden))
        del worksheet.column_dimensions[letter]
    for index, width, hidden in pending:
        target = worksheet.column_dimensions[get_column_letter(index + delta)]
        target.width = width
        target.hidden = hidden


def _parse_coordinate(key):
    parts = str(key).split(',')
    if len(parts) != 2:
        raise XlsxPatchExportError(f'Bad cell coordinate {key!r}.')
    try:
        row_pos = int(parts[0])
        col_pos = int(parts[1])
    except ValueError as exc:
        raise XlsxPatchExportError(f'Bad cell coordinate {key!r}.') from exc
    if row_pos < 0 or col_pos < 0:
        raise XlsxPatchExportError(f'Bad cell coordinate {key!r}.')
    return row_pos, col_pos


def _active_at(model, sheet, position, label):
    obj = model.objects.filter(sheet=sheet, position=position, is_deleted=False).first()
    if obj is None:
        raise XlsxPatchExportError(f'No active {label} at position {position}.')
    return obj


def _live_cell(sheet, row, column):
    return (
        Cell.objects.filter(sheet=sheet, row=row, column=column, is_deleted=False)
        .first()
    )


def _excel_cell(worksheet, key):
    row_pos, col_pos = _parse_coordinate(key)
    return worksheet.cell(row=row_pos + 1, column=col_pos + 1)


def _reject_sparkline_changes(sheet, cell_changes):
    for key, change in cell_changes.items():
        raw = (change or {}).get('raw_input')
        if raw is None:
            continue
        if is_sparkline(str(raw)):
            raise XlsxPatchExportError(f'Cell {key} is a sparkline and cannot be patched.')
        row_pos, col_pos = _parse_coordinate(key)
        row = _active_at(SheetRow, sheet, row_pos, 'row')
        column = _active_at(SheetColumn, sheet, col_pos, 'column')
        cell = _live_cell(sheet, row, column)
        if cell is not None and is_sparkline(cell.raw_input):
            raise XlsxPatchExportError(f'Cell {key} is a sparkline and cannot be patched.')


def _patch_value(worksheet, sheet, key, change):
    row_pos, col_pos = _parse_coordinate(key)
    excel_cell = worksheet.cell(row=row_pos + 1, column=col_pos + 1)
    raw = (change or {}).get('raw_input')
    if raw is None:
        excel_cell.value = None
        return

    row = _active_at(SheetRow, sheet, row_pos, 'row')
    column = _active_at(SheetColumn, sheet, col_pos, 'column')
    cell = _live_cell(sheet, row, column)
    if cell is None:
        raise XlsxPatchExportError(f'No active cell at {key} to convert for export.')

    stored = (cell.raw_input or '').strip()
    if stored.startswith('='):
        excel_cell.value = stored
        return
    excel_cell.value = _cell_export_value(cell)


def _apply_font_family(font, value) -> bool:
    """Return True when the font name was changed."""
    if value is None:
        font.name = _OPENPYXL_DEFAULT_FONT_NAME
        font.scheme = None
        return True
    mapped = _EXCEL_FONT_NAMES.get(str(value).strip())
    if not mapped:
        return False
    font.name = mapped
    font.scheme = None
    return True


def _patch_font(excel_cell, changes):
    if not any(field in changes for field in _FONT_FIELDS):
        return
    font = copy(excel_cell.font)
    changed = False
    if 'bold' in changes:
        font.bold = bool(changes['bold'])
        changed = True
    if 'italic' in changes:
        font.italic = bool(changes['italic'])
        changed = True
    if 'strikethrough' in changes:
        font.strike = bool(changes['strikethrough'])
        changed = True
    if 'font_family' in changes and _apply_font_family(font, changes['font_family']):
        changed = True
    if 'font_size' in changes:
        font.size = (
            _OPENPYXL_DEFAULT_FONT_SIZE
            if changes['font_size'] is None
            else changes['font_size']
        )
        changed = True
    if 'text_color' in changes:
        if changes['text_color'] is None:
            font.color = _OPENPYXL_DEFAULT_FONT_COLOR
        else:
            rgb = _excel_rgb(changes['text_color'])
            if not rgb:
                raise XlsxPatchExportError(
                    f'Unsupported text color {changes["text_color"]!r}.'
                )
            font.color = rgb
        changed = True
    if changed:
        excel_cell.font = font


def _patch_number_format(excel_cell, changes):
    if 'number_format' not in changes:
        return
    value = changes['number_format']
    if value is None:
        excel_cell.number_format = 'General'
        return
    mapped = _excel_number_format(value)
    excel_cell.number_format = mapped or 'General'


def _patch_format(worksheet, key, changes):
    excel_cell = _excel_cell(worksheet, key)
    _patch_font(excel_cell, changes)
    _patch_number_format(excel_cell, changes)


def _patch_highlight(worksheet, key, color):
    excel_cell = _excel_cell(worksheet, key)
    if color is None:
        excel_cell.fill = PatternFill(fill_type=None)
        return
    if not _excel_rgb(color):
        raise XlsxPatchExportError(f'Unsupported highlight color {color!r}.')
    _apply_fill(excel_cell, color)