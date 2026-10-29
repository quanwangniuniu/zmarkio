import io
import tempfile
from contextlib import contextmanager
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from rest_framework.test import APIClient

from spreadsheet.import_baseline import build_import_baseline
from spreadsheet.import_diff import diff_sheet
from spreadsheet.models import (
    Cell,
    CellValueType,
    SheetColumn,
    SheetImportSource,
    SheetRow,
    SpreadsheetCellFormat,
    SpreadsheetHighlight,
)
from spreadsheet.services import SheetService
from spreadsheet.tests.test_services import (
    create_test_organization,
    create_test_project,
    create_test_sheet,
    create_test_spreadsheet,
    create_test_user,
)
from spreadsheet.xlsx_export import build_sheet_workbook
from spreadsheet.xlsx_patch_export import XlsxPatchExportError, build_patched_workbook

def _workbook_bytes(build):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Sheet1'
    build(workbook, worksheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@contextmanager
def _open_exported(content):
    workbook = load_workbook(io.BytesIO(content))
    try:
        yield workbook
    finally:
        workbook.close()


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class TestXlsxPatchExport(TestCase):
    def setUp(self):
        self.user = create_test_user()
        self.organization = create_test_organization()
        self.project = create_test_project(self.organization, owner=self.user)
        self.spreadsheet = create_test_spreadsheet(self.project)
        self.sheet = create_test_sheet(self.spreadsheet)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.url = (
            f'/api/spreadsheet/spreadsheets/{self.spreadsheet.slug}/'
            f'sheets/{self.sheet.id}/export.xlsx'
        )

    def _cell(self, row_pos, col_pos, raw):
        row, _created = SheetRow.objects.get_or_create(sheet=self.sheet, position=row_pos)
        column, _created = SheetColumn.objects.get_or_create(
            sheet=self.sheet,
            position=col_pos,
            defaults={'name': chr(ord('A') + col_pos)},
        )
        return Cell.objects.create(
            sheet=self.sheet,
            row=row,
            column=column,
            raw_input=raw,
            string_value=raw,
            value_type=CellValueType.STRING,
        )

    def _save_excel_source(self, original_bytes, worksheet_name='Sheet1'):
        """Call only after the sheet matches the workbook and before any edit."""
        return SheetImportSource.objects.create(
            sheet=self.sheet,
            origin=SheetImportSource.Origin.EXCEL,
            worksheet_name=worksheet_name,
            original_filename='book.xlsx',
            baseline=build_import_baseline(self.sheet),
            file=ContentFile(original_bytes, name='book.xlsx'),
        )

    def _assert_rebuild(self):
        captured = {}
        real = build_sheet_workbook

        def wrapper(sheet):
            captured['content'] = real(sheet)
            return captured['content']

        with patch('spreadsheet.views.build_sheet_workbook', side_effect=wrapper):
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, captured['content'])

    def test_no_source_uses_rebuild(self):
        self._cell(0, 0, 'plain')
        self._assert_rebuild()

    def test_google_source_uses_rebuild(self):
        self._cell(0, 0, 'plain')
        SheetImportSource.objects.create(
            sheet=self.sheet,
            origin=SheetImportSource.Origin.GOOGLE,
            worksheet_name='Budget',
            google_spreadsheet_id='abc',
            google_sheet_id=0,
        )
        self._assert_rebuild()

    def test_no_changes_return_original_bytes(self):
        self._cell(0, 0, 'same')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'same'))
        self._save_excel_source(original)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, original)

    def test_value_change_patches_only_that_cell(self):
        changed = self._cell(0, 0, 'old')
        self._cell(0, 1, 'keep')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'old'),
                worksheet.__setitem__('B1', 'keep'),
            )
        )
        self._save_excel_source(original)
        changed.raw_input = 'new'
        changed.string_value = 'new'
        changed.save(update_fields=['raw_input', 'string_value'])
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'new')
            self.assertEqual(worksheet['B1'].value, 'keep')

    def test_bold_change_keeps_underline(self):
        self._cell(0, 0, 'styled')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'styled'),
                setattr(worksheet['A1'], 'font', Font(bold=False, underline='single')),
            )
        )
        self._save_excel_source(original)
        SpreadsheetCellFormat.objects.create(
            sheet=self.sheet, row_index=0, column_index=0, bold=True,
        )
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.font.underline, 'single')

    def test_highlight_change_keeps_border_and_alignment(self):
        self._cell(0, 0, 'boxed')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'boxed'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                setattr(worksheet['A1'], 'border', Border(left=Side(style='thin'))),
                setattr(worksheet['A1'], 'alignment', Alignment(horizontal='center')),
            )
        )
        self._save_excel_source(original)
        SpreadsheetHighlight.objects.create(
            spreadsheet=self.spreadsheet,
            sheet=self.sheet,
            scope='CELL',
            row_index=0,
            col_index=0,
            color='#BFDBFE',
        )
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertIn('BFDBFE', cell.fill.fgColor.rgb)
            self.assertEqual(cell.font.underline, 'single')
            self.assertEqual(cell.border.left.style, 'thin')
            self.assertEqual(cell.alignment.horizontal, 'center')

    def test_unchanged_custom_number_format_survives_bold(self):
        self._cell(0, 0, '10')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 10),
                setattr(worksheet['A1'], 'number_format', '[$AUD]#,##0.00'),
            )
        )
        self._save_excel_source(original)
        SpreadsheetCellFormat.objects.create(
            sheet=self.sheet, row_index=0, column_index=0, bold=True,
        )
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.number_format, '[$AUD]#,##0.00')

    def test_other_worksheet_is_preserved(self):
        self._cell(0, 0, 'old')

        def build(workbook, worksheet):
            worksheet['A1'] = 'old'
            other = workbook.create_sheet('Sheet2')
            other['A1'] = 'other-sheet'

        original = _workbook_bytes(build)
        source_cell = Cell.objects.get(sheet=self.sheet)
        self._save_excel_source(original)
        source_cell.raw_input = 'new'
        source_cell.string_value = 'new'
        source_cell.save(update_fields=['raw_input', 'string_value'])
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook.sheetnames, ['Sheet1', 'Sheet2'])
            self.assertEqual(workbook['Sheet1']['A1'].value, 'new')
            self.assertEqual(workbook['Sheet2']['A1'].value, 'other-sheet')

    def test_patch_failure_falls_back_to_rebuild(self):
        cell = self._cell(0, 0, 'old')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'old'))
        self._save_excel_source(original)
        cell.raw_input = 'new'
        cell.string_value = 'new'
        cell.save(update_fields=['raw_input', 'string_value'])
        with patch(
            'spreadsheet.views.build_patched_workbook',
            side_effect=XlsxPatchExportError('patch failed'),
        ):
            self._assert_rebuild()

    def test_source_read_failure_falls_back_to_rebuild(self):
        self._cell(0, 0, 'same')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'same'))
        self._save_excel_source(original)
        with patch(
            'django.db.models.fields.files.FieldFile.open',
            side_effect=OSError('unreadable'),
        ):
            self._assert_rebuild()

    def test_column_insert_only_downloads_the_shifted_sheet(self):
        self._cell(0, 0, 'left')
        self._cell(0, 1, 'right')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                worksheet.__setitem__('B1', 'right'),
                setattr(worksheet['B1'], 'font', Font(underline='single')),
            )
        )
        source = self._save_excel_source(original)
        with source.file.open('rb') as handle:
            stored = handle.read()
        SheetService.insert_columns(self.sheet, position=1, count=1)

        called = {}
        real_patch = build_patched_workbook

        def wrapper(source_arg, diff):
            called['patch'] = True
            return real_patch(source_arg, diff)

        with patch('spreadsheet.views.build_patched_workbook', side_effect=wrapper):
            with patch('spreadsheet.views.build_sheet_workbook') as rebuild:
                response = self.client.get(self.url)

        self.assertEqual(response.status_code, 200)
        self.assertTrue(called.get('patch'))
        rebuild.assert_not_called()
        self.assertNotEqual(response.content, original)
        with source.file.open('rb') as handle:
            self.assertEqual(handle.read(), stored)
        with _open_exported(response.content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'left')
            self.assertIsNone(worksheet['B1'].value)
            self.assertEqual(worksheet['C1'].value, 'right')
            self.assertEqual(worksheet['C1'].font.underline, 'single')


    def _structure_diff(self, operations, cells=None):
        return {
            'structural_operations': operations,
            'cells': cells or {},
            'formats': {},
            'highlights': {'cells': {}},
        }

    def test_column_insert_moves_the_underlined_cell(self):
        self._cell(0, 0, 'left')
        self._cell(0, 1, 'right')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                worksheet.__setitem__('B1', 'right'),
                setattr(worksheet['B1'], 'font', Font(underline='single')),
            )
        )
        source = self._save_excel_source(original)
        with source.file.open('rb') as handle:
            stored = handle.read()
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'COL_INSERT', 'anchor_position': 1, 'count': 1},
        ]))
        with source.file.open('rb') as handle:
            self.assertEqual(handle.read(), stored)
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'left')
            self.assertIsNone(workbook['Sheet1']['B1'].value)
            self.assertEqual(workbook['Sheet1']['C1'].value, 'right')
            self.assertEqual(workbook['Sheet1']['C1'].font.underline, 'single')

    def test_row_insert_moves_the_cell_down(self):
        self._cell(0, 0, 'top')
        self._cell(1, 0, 'bottom')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.__setitem__('A2', 'bottom'),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_INSERT', 'anchor_position': 1, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'top')
            self.assertIsNone(workbook['Sheet1']['A2'].value)
            self.assertEqual(workbook['Sheet1']['A3'].value, 'bottom')

    def test_row_delete_shifts_the_following_cell_up(self):
        self._cell(0, 0, 'drop')
        self._cell(1, 0, 'keep')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'drop'),
                worksheet.__setitem__('A2', 'keep'),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_DELETE', 'anchor_position': 0, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'keep')

    def test_row_then_column_insert_applies_in_order(self):
        self._cell(0, 0, 'A')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'A'))
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_INSERT', 'anchor_position': 0, 'count': 1},
            {'op_type': 'COL_INSERT', 'anchor_position': 0, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            self.assertIsNone(workbook['Sheet1']['A1'].value)
            self.assertEqual(workbook['Sheet1']['B2'].value, 'A')

    def test_final_coordinate_patch_is_applied_after_insert(self):
        cell = self._cell(0, 1, 'right')
        self._cell(0, 0, 'left')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                worksheet.__setitem__('B1', 'right'),
                setattr(worksheet['B1'], 'font', Font(underline='single')),
            )
        )
        source = self._save_excel_source(original)
        SheetService.insert_columns(self.sheet, position=1, count=1)
        cell.raw_input = 'edited'
        cell.string_value = 'edited'
        cell.save(update_fields=['raw_input', 'string_value'])
        content = build_patched_workbook(source, diff_sheet(source))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['C1'].value, 'edited')
            self.assertEqual(workbook['Sheet1']['C1'].font.underline, 'single')

    def test_merge_starting_at_the_insert_moves(self):
        self._cell(0, 0, 'top')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.merge_cells('A1:A2'),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_INSERT', 'anchor_position': 0, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A2'].value, 'top')
            self.assertIn('A2:A3', {str(item) for item in worksheet.merged_cells.ranges})

    def test_merge_across_the_insert_raises(self):
        self._cell(0, 0, 'top')
        self._cell(1, 0, 'mid')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.__setitem__('A2', 'mid'),
                worksheet.merge_cells('A1:A3'),
            )
        )
        source = self._save_excel_source(original)
        with source.file.open('rb') as handle:
            stored = handle.read()
        with self.assertRaises(XlsxPatchExportError):
            build_patched_workbook(source, self._structure_diff([
                {'op_type': 'ROW_INSERT', 'anchor_position': 1, 'count': 1},
            ]))
        with source.file.open('rb') as handle:
            self.assertEqual(handle.read(), stored)

    def test_other_worksheet_is_not_shifted(self):
        self._cell(0, 0, 'A')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'A'),
                workbook.create_sheet('Sheet2').__setitem__('A1', 'other'),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_INSERT', 'anchor_position': 0, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A2'].value, 'A')
            self.assertEqual(workbook['Sheet2']['A1'].value, 'other')

    def test_column_delete_shifts_the_following_cell_left(self):
        self._cell(0, 0, 'drop')
        self._cell(0, 1, 'keep')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'drop'),
                worksheet.__setitem__('B1', 'keep'),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'COL_DELETE', 'anchor_position': 0, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'keep')

    def test_column_width_at_the_insert_moves(self):
        self._cell(0, 0, 'left')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                setattr(worksheet.column_dimensions['B'], 'width', 22),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'COL_INSERT', 'anchor_position': 1, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'left')
            self.assertEqual(worksheet.column_dimensions['C'].width, 22)
            self.assertNotIn('B', worksheet.column_dimensions)

    def test_column_width_before_the_insert_is_allowed(self):
        self._cell(0, 0, 'left')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                setattr(worksheet.column_dimensions['A'], 'width', 22),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'COL_INSERT', 'anchor_position': 2, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'left')
            self.assertEqual(workbook['Sheet1'].column_dimensions['A'].width, 22)

    def test_row_height_at_the_insert_moves(self):
        self._cell(0, 0, 'top')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                setattr(worksheet.row_dimensions[2], 'height', 30),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_INSERT', 'anchor_position': 1, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'top')
            self.assertEqual(worksheet.row_dimensions[3].height, 30)
            self.assertNotIn(2, worksheet.row_dimensions)

    def test_relative_formula_is_patched_after_insert(self):
        formula = self._cell(0, 0, '=B1')
        self._cell(0, 1, '2')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', '=B1'),
                worksheet.__setitem__('B1', 2),
            )
        )
        source = self._save_excel_source(original)
        SheetService.insert_columns(self.sheet, position=1, count=1)
        formula.refresh_from_db()
        self.assertEqual(formula.raw_input, '=C1')
        content = build_patched_workbook(source, diff_sheet(source))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, '=C1')
            self.assertEqual(workbook['Sheet1']['C1'].value, 2)

    def test_absolute_formula_refuses_structural_replay(self):
        self._cell(0, 0, '=$B$1')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', '=$B$1'))
        source = self._save_excel_source(original)
        with self.assertRaises(XlsxPatchExportError):
            build_patched_workbook(source, self._structure_diff([
                {'op_type': 'COL_INSERT', 'anchor_position': 1, 'count': 1},
            ]))

    def test_value_edit_keeps_a_merge_when_structure_is_unchanged(self):
        cell = self._cell(0, 0, 'old')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'old'),
                worksheet.merge_cells('B1:C1'),
                setattr(worksheet.column_dimensions['B'], 'width', 18),
            )
        )
        source = self._save_excel_source(original)
        cell.raw_input = 'new'
        cell.string_value = 'new'
        cell.save(update_fields=['raw_input', 'string_value'])
        content = build_patched_workbook(source, diff_sheet(source))
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'new')
            self.assertIn('B1:C1', {str(item) for item in workbook['Sheet1'].merged_cells.ranges})
            self.assertEqual(workbook['Sheet1'].column_dimensions['B'].width, 18)

    def test_delete_last_row_keeps_underline_on_the_rows_above(self):
        self._cell(0, 0, 'keep')
        self._cell(1, 0, 'keep2')
        self._cell(2, 0, 'drop')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'keep'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                worksheet.__setitem__('A2', 'keep2'),
                worksheet.__setitem__('A3', 'drop'),
                setattr(worksheet.row_dimensions[3], 'height', 30),
            )
        )
        source = self._save_excel_source(original)
        with source.file.open('rb') as handle:
            stored = handle.read()
        SheetService.delete_rows(self.sheet, position=2, count=1)
        with patch('spreadsheet.views.build_sheet_workbook') as rebuild:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        rebuild.assert_not_called()
        with source.file.open('rb') as handle:
            self.assertEqual(handle.read(), stored)
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'keep')
            self.assertEqual(workbook['Sheet1']['A1'].font.underline, 'single')
            self.assertEqual(workbook['Sheet1']['A2'].value, 'keep2')
            self.assertIsNone(workbook['Sheet1']['A3'].value)

    def test_delete_middle_row_moves_the_underlined_cell_up(self):
        self._cell(0, 0, 'top')
        self._cell(1, 0, 'drop')
        self._cell(2, 0, 'bottom')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.__setitem__('A2', 'drop'),
                setattr(worksheet.row_dimensions[2], 'height', 30),
                worksheet.__setitem__('A3', 'bottom'),
                setattr(worksheet['A3'], 'font', Font(underline='single')),
            )
        )
        source = self._save_excel_source(original)
        SheetService.delete_rows(self.sheet, position=1, count=1)
        with patch('spreadsheet.views.build_sheet_workbook') as rebuild:
            response = self.client.get(self.url)
        rebuild.assert_not_called()
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'top')
            self.assertEqual(workbook['Sheet1']['A2'].value, 'bottom')
            self.assertEqual(workbook['Sheet1']['A2'].font.underline, 'single')

    def test_delete_column_keeps_underline_on_the_next_column(self):
        self._cell(0, 0, 'drop')
        self._cell(0, 1, 'keep')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'drop'),
                setattr(worksheet.column_dimensions['A'], 'width', 22),
                worksheet.__setitem__('B1', 'keep'),
                setattr(worksheet['B1'], 'font', Font(underline='single')),
            )
        )
        source = self._save_excel_source(original)
        SheetService.delete_columns(self.sheet, position=0, count=1)
        with patch('spreadsheet.views.build_sheet_workbook') as rebuild:
            response = self.client.get(self.url)
        rebuild.assert_not_called()
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'keep')
            self.assertEqual(workbook['Sheet1']['A1'].font.underline, 'single')

    def test_height_below_a_deleted_row_stays_in_the_patched_workbook(self):
        self._cell(0, 0, 'top')
        self._cell(1, 0, 'drop')
        self._cell(2, 0, 'bottom')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.__setitem__('A2', 'drop'),
                worksheet.__setitem__('A3', 'bottom'),
                setattr(worksheet['A3'], 'font', Font(underline='single')),
                setattr(worksheet.row_dimensions[3], 'height', 30),
            )
        )
        self._save_excel_source(original)
        SheetService.delete_rows(self.sheet, position=1, count=1)
        with patch('spreadsheet.views.build_sheet_workbook') as rebuild:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        rebuild.assert_not_called()
        with _open_exported(response.content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'top')
            self.assertEqual(worksheet['A2'].value, 'bottom')
            self.assertEqual(worksheet['A2'].font.underline, 'single')
            self.assertEqual(worksheet.row_dimensions[2].height, 30)
        
    def test_delete_last_row_removes_its_height(self):
        self._cell(0, 0, 'keep')
        self._cell(2, 0, 'drop')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'keep'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                worksheet.__setitem__('A3', 'drop'),
                setattr(worksheet.row_dimensions[3], 'height', 30),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_DELETE', 'anchor_position': 2, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertNotIn(3, worksheet.row_dimensions)
            self.assertEqual(worksheet['A1'].font.underline, 'single')

    def test_delete_column_removes_its_width(self):
        self._cell(0, 0, 'drop')
        self._cell(0, 1, 'keep')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'drop'),
                setattr(worksheet.column_dimensions['A'], 'width', 22),
                worksheet.__setitem__('B1', 'keep'),
                setattr(worksheet['B1'], 'font', Font(underline='single')),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'COL_DELETE', 'anchor_position': 0, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertNotIn('A', worksheet.column_dimensions)
            self.assertEqual(worksheet['A1'].value, 'keep')
            self.assertEqual(worksheet['A1'].font.underline, 'single')

    def test_delete_row_removes_a_merge_on_that_row(self):
        self._cell(0, 0, 'keep')
        self._cell(2, 0, 'drop')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'keep'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                worksheet.__setitem__('A3', 'drop'),
                worksheet.merge_cells('A3:B3'),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_DELETE', 'anchor_position': 2, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertNotIn('A3:B3', {str(item) for item in worksheet.merged_cells.ranges})
            self.assertEqual(worksheet['A1'].font.underline, 'single')

    def test_second_delete_does_not_see_the_removed_height(self):
        self._cell(0, 0, 'keep')
        self._cell(3, 0, 'drop')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'keep'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                worksheet.__setitem__('A4', 'drop'),
                setattr(worksheet.row_dimensions[4], 'height', 30),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_DELETE', 'anchor_position': 3, 'count': 1},
            {'op_type': 'ROW_DELETE', 'anchor_position': 2, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertNotIn(4, worksheet.row_dimensions)
            self.assertEqual(worksheet['A1'].font.underline, 'single')

    def test_height_below_a_deleted_row_moves_with_the_row(self):
        self._cell(0, 0, 'top')
        self._cell(2, 0, 'bottom')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.__setitem__('A2', 'drop'),
                worksheet.__setitem__('A3', 'bottom'),
                setattr(worksheet.row_dimensions[3], 'height', 30),
            )
        )
        source = self._save_excel_source(original)
        content = build_patched_workbook(source, self._structure_diff([
            {'op_type': 'ROW_DELETE', 'anchor_position': 1, 'count': 1},
        ]))
        with _open_exported(content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'top')
            self.assertEqual(worksheet['A2'].value, 'bottom')
            self.assertEqual(worksheet.row_dimensions[2].height, 30)
            self.assertNotIn(3, worksheet.row_dimensions)
        
    def test_delete_column_with_trailing_column_keeps_underline(self):
        self._cell(0, 0, 'keep')
        self._cell(0, 1, 'drop')
        self._cell(0, 2, 'next')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'keep'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                worksheet.__setitem__('B1', 'drop'),
                worksheet.__setitem__('C1', 'next'),
            )
        )
        source = self._save_excel_source(original)
        SheetService.delete_columns(self.sheet, position=1, count=1)
        SheetColumn.objects.create(sheet=self.sheet, position=2, name='pad')
        with patch('spreadsheet.views.build_sheet_workbook') as rebuild:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        rebuild.assert_not_called()
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'keep')
            self.assertEqual(workbook['Sheet1']['A1'].font.underline, 'single')
            self.assertEqual(workbook['Sheet1']['B1'].value, 'next')

    def _save_google_source(self):
        return SheetImportSource.objects.create(
            sheet=self.sheet,
            origin=SheetImportSource.Origin.GOOGLE,
            worksheet_name='Budget',
            google_spreadsheet_id='google-book',
            google_sheet_id=0,
            baseline=build_import_baseline(self.sheet),
        )

    def _connect_google(self):
        from google_docs_integration.models import GoogleDocsConnection

        connection = GoogleDocsConnection(user=self.user, is_active=True)
        connection.set_access_token('token-abc')
        connection.save()
        return connection

    def test_google_download_without_edits_returns_the_exported_workbook(self):
        self._cell(0, 0, 'same')
        self._save_google_source()
        self._connect_google()
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'same'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                workbook.create_sheet('Other'),
            )
        )
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            return_value=(original, 'Sheet1'),
        ) as fetch:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, original)
        fetch.assert_called_once()

    def test_google_download_patches_a_value_and_keeps_underline(self):
        changed = self._cell(0, 0, 'old')
        self._save_google_source()
        self._connect_google()
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'old'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
            )
        )
        changed.raw_input = 'new'
        changed.string_value = 'new'
        changed.save(update_fields=['raw_input', 'string_value'])
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            return_value=(original, 'Sheet1'),
        ):
            response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'new')
            self.assertEqual(workbook['Sheet1']['A1'].font.underline, 'single')

    def test_google_highlight_keeps_border_and_alignment(self):
        self._cell(0, 0, 'boxed')
        self._save_google_source()
        self._connect_google()
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'boxed'),
                setattr(worksheet['A1'], 'border', Border(left=Side(style='thin'))),
                setattr(worksheet['A1'], 'alignment', Alignment(horizontal='center')),
            )
        )
        SpreadsheetHighlight.objects.create(
            spreadsheet=self.spreadsheet,
            sheet=self.sheet,
            scope='CELL',
            row_index=0,
            col_index=0,
            color='#BFDBFE',
        )
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            return_value=(original, 'Sheet1'),
        ):
            response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertIn('BFDBFE', cell.fill.fgColor.rgb)
            self.assertEqual(cell.border.left.style, 'thin')
            self.assertEqual(cell.alignment.horizontal, 'center')

    def test_google_column_insert_moves_the_underlined_cell(self):
        self._cell(0, 0, 'left')
        self._cell(0, 1, 'right')
        self._save_google_source()
        self._connect_google()
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                worksheet.__setitem__('B1', 'right'),
                setattr(worksheet['B1'], 'font', Font(underline='single')),
            )
        )
        SheetService.insert_columns(self.sheet, position=1, count=1, created_by=self.user)
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            return_value=(original, 'Sheet1'),
        ):
            response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'left')
            self.assertIsNone(worksheet['B1'].value)
            self.assertEqual(worksheet['C1'].value, 'right')
            self.assertEqual(worksheet['C1'].font.underline, 'single')

    def test_google_row_delete_shifts_the_underlined_cell(self):
        self._cell(0, 0, 'drop')
        self._cell(1, 0, 'keep')
        self._save_google_source()
        self._connect_google()
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'drop'),
                worksheet.__setitem__('A2', 'keep'),
                setattr(worksheet['A2'], 'font', Font(underline='single')),
            )
        )
        SheetService.delete_rows(self.sheet, position=0, count=1, created_by=self.user)
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            return_value=(original, 'Sheet1'),
        ):
            response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'keep')
            self.assertEqual(worksheet['A1'].font.underline, 'single')

    def test_google_undo_returns_the_exported_workbook(self):
        from spreadsheet.models import SheetStructureOperation

        self._cell(0, 0, 'top')
        self._cell(1, 0, 'bottom')
        self._save_google_source()
        self._connect_google()
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'top'),
                worksheet.__setitem__('A2', 'bottom'),
            )
        )
        deleted = SheetService.delete_rows(self.sheet, position=1, count=1, created_by=self.user)
        operation = SheetStructureOperation.objects.get(id=deleted['operation_id'])
        SheetService.revert_structure_operation(self.sheet, operation)
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            return_value=(original, 'Sheet1'),
        ):
            response = self.client.get(self.url)
        self.assertEqual(response.content, original)

    def test_missing_google_worksheet_rebuilds(self):
        self._cell(0, 0, 'plain')
        self._save_google_source()
        self._connect_google()
        with patch(
            'spreadsheet.views.export_google_spreadsheet_xlsx',
            side_effect=ValueError('google worksheet was not found'),
        ):
            self._assert_rebuild()

    def test_workbook_bytes_use_the_live_worksheet_name(self):
        changed = self._cell(0, 0, 'old')
        source = self._save_google_source()
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'old'))
        changed.raw_input = 'new'
        changed.string_value = 'new'
        changed.save(update_fields=['raw_input', 'string_value'])
        content = build_patched_workbook(
            source,
            diff_sheet(source),
            workbook_bytes=original,
            worksheet_name='Sheet1',
        )
        with _open_exported(content) as workbook:
            self.assertEqual(workbook['Sheet1']['A1'].value, 'new')
        with self.assertRaises(XlsxPatchExportError):
            build_patched_workbook(
                source,
                diff_sheet(source),
                workbook_bytes=original,
                worksheet_name='Missing',
            )